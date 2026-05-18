from __future__ import annotations

import csv
import io
from pathlib import Path
from zipfile import BadZipFile

from bs4 import BeautifulSoup

from agentmail.parser import html_to_text


TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".tsv", ".html", ".htm"}
MAX_TEXT_BYTES = 512_000


def extract_text(filename: str, content_type: str | None, payload: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md"} or (content_type or "").startswith("text/plain"):
        return payload[:MAX_TEXT_BYTES].decode("utf-8", errors="replace")
    if suffix in {".html", ".htm"} or content_type == "text/html":
        html = payload[:MAX_TEXT_BYTES].decode("utf-8", errors="replace")
        return html_to_text(html)
    if suffix in {".csv", ".tsv"}:
        delimiter = "\t" if suffix == ".tsv" else ","
        text = payload[:MAX_TEXT_BYTES].decode("utf-8", errors="replace")
        rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
        return "\n".join(" | ".join(row) for row in rows[:200])
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(payload))
            return "\n".join((page.extract_text() or "") for page in reader.pages[:20]).strip()
        except Exception:
            return ""
    if suffix == ".docx":
        try:
            import docx

            document = docx.Document(io.BytesIO(payload))
            return "\n".join(paragraph.text for paragraph in document.paragraphs).strip()
        except (BadZipFile, Exception):
            return ""
    if suffix == ".xlsx":
        try:
            import openpyxl

            workbook = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
            lines: list[str] = []
            for sheet in workbook.worksheets[:5]:
                lines.append(f"Sheet: {sheet.title}")
                for row in sheet.iter_rows(max_row=50, values_only=True):
                    values = ["" if value is None else str(value) for value in row]
                    lines.append(" | ".join(values))
            return "\n".join(lines).strip()
        except Exception:
            return ""
    return BeautifulSoup(payload[:MAX_TEXT_BYTES].decode("utf-8", errors="ignore"), "html.parser").get_text("\n").strip()
